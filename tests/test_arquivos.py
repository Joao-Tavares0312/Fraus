"""Testes da leitura de arquivo: csv, xlsx, docx e pdf.

Metade daqui verifica as MENSAGENS DE ERRO, e não é excesso: arquivo que não
entra é o caso comum, porque cada operação exporta do sistema que tem, não do
formato que o Fraus pede. Recusa que não diz o que se esperava obriga a pessoa
a adivinhar, e é o defeito que originou este módulo.
"""

import io

import pytest

from fraus.ingest.arquivos import ArquivoIlegivelError, extrair

CABECALHO = "conversa_id,canal,autor,texto,enviada_em,escalou_para_humano"
CSV_FRAUS = (
    f"{CABECALHO}\n"
    "c1,webchat,cliente,meu pedido nao chegou,2026-05-14T10:00:00+00:00,false\n"
    "c1,webchat,bot,vou verificar,2026-05-14T10:00:12+00:00,false\n"
)

TRANSCRICAO = (
    "Cliente: Bom dia, minha cobranca veio duplicada\n"
    "Bot: Entendi! Vou verificar seu cadastro.\n"
    "Cliente: Ok obrigado\n"
    "Atendente: Localizei a duplicidade, ja estornei.\n"
)


def _docx(texto: str) -> bytes:
    import docx

    documento = docx.Document()
    for linha in texto.splitlines():
        documento.add_paragraph(linha)
    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def _xlsx(linhas: list[list]) -> bytes:
    from openpyxl import Workbook

    pasta = Workbook()
    for linha in linhas:
        pasta.active.append(linha)
    buffer = io.BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# formatos que trazem horario: recebem nota


def test_csv_no_formato_do_fraus(cliente_nao_usado=None):
    extracao = extrair("conversa.csv", CSV_FRAUS.encode("utf-8"))
    assert len(extracao.conversas) == 1
    assert extracao.tem_tempo is True
    assert "Fraus" in extracao.formato


def test_planilha_com_as_mesmas_colunas_do_csv():
    dados = _xlsx([
        CABECALHO.split(","),
        ["c1", "webchat", "cliente", "meu pedido nao chegou", "2026-05-14T10:00:00+00:00", "false"],
        ["c1", "webchat", "bot", "vou verificar", "2026-05-14T10:00:12+00:00", "false"],
    ])
    extracao = extrair("planilha.xlsx", dados)
    assert len(extracao.conversas[0].mensagens) == 2
    assert extracao.tem_tempo is True


def test_planilha_da_totalk_com_data_de_celula():
    # O export .xlsx da Totalk traz a data como DATETIME de celula, que o
    # leitor de planilha converte para ISO -- nao como o MM/DD/YYYY do export
    # .csv. Rejeitar a ISO fazia o arquivo INTEIRO voltar como "nenhuma
    # conversa valida" (31/08/2026, export real do Joao): o adaptador era
    # acionado, e todas as linhas morriam em "data invalida".
    from datetime import datetime

    url = "https://app.totalk.chat/redirect?type=SE&id=0bf841a6-e15f-446c-96cf-391c361cfb17"
    cabecalho = [
        "Conta/Nome", "Canal/Plataforma", "Contato/Nome",
        "Mensagem/Data de criação", "Mensagem/Quem enviou",
        "Mensagem/Conteúdo", "Conversa",
    ]
    dados = _xlsx([
        cabecalho,
        ["Empresa X", "WhatsApp", "Cliente Y",
         datetime(2025, 10, 6, 20, 2, 53), "De: Cliente Y Para: Empresa X",
         "meu pedido nao chegou", url],
        ["Empresa X", "WhatsApp", "Cliente Y",
         datetime(2025, 10, 6, 20, 3, 12), "De: Empresa X Para: Cliente Y",
         "vou verificar", url],
    ])
    extracao = extrair("message-export.xlsx", dados)
    assert extracao.formato == "export da Totalk"
    assert len(extracao.conversas) == 1
    assert len(extracao.conversas[0].mensagens) == 2
    # O fuso continua o da Totalk (-03:00): a data de celula vem naive.
    assert extracao.conversas[0].mensagens[0].enviada_em.utcoffset() is not None


def test_csv_da_totalk_continua_com_a_data_americana():
    # A armadilha 1 nao mudou de lugar: no export .csv a data segue MM/DD/YYYY,
    # e aceitar ISO nao pode quebrar o caminho antigo.
    url = "https://app.totalk.chat/redirect?type=SE&id=0bf841a6-e15f-446c-96cf-391c361cfb17"
    csv_totalk = (
        "Conta/Nome,Canal/Plataforma,Mensagem/Data de criação,"
        "Mensagem/Quem enviou,Mensagem/Conteúdo,Conversa\n"
        f'Empresa X,WhatsApp,10/06/2025 20:02:53,De: Cliente Y Para: Empresa X,meu pedido nao chegou,{url}\n'
        f'Empresa X,WhatsApp,10/06/2025 20:03:12,De: Empresa X Para: Cliente Y,vou verificar,{url}\n'
    )
    extracao = extrair("export.csv", csv_totalk.encode("utf-8"))
    assert extracao.formato == "export da Totalk"
    assert extracao.conversas[0].mensagens[0].enviada_em.month == 10
    assert extracao.conversas[0].mensagens[0].enviada_em.day == 6


def test_csv_em_latin1_nao_e_recusado_por_causa_de_acento():
    """Export de Windows sai em latin-1. Recusar por isso seria perder o arquivo."""
    texto = CSV_FRAUS.replace("meu pedido nao chegou", "meu pedido não chegou")
    extracao = extrair("conversa.csv", texto.encode("latin-1"))
    assert "não chegou" in extracao.conversas[0].mensagens[0].texto


# ---------------------------------------------------------------------------
# transcricao em prosa: SEM horario, SEM nota


def test_docx_vira_transcricao_com_os_autores_reconhecidos():
    extracao = extrair("conversa.docx", _docx(TRANSCRICAO))
    autores = [m.autor for m in extracao.conversas[0].mensagens]
    assert autores == ["cliente", "bot", "cliente", "humano"]


def test_transcricao_sem_horario_marca_tem_tempo_falso_e_avisa():
    """O aviso e a razao de a nota nao sair -- sem ele, a ausencia parece bug."""
    extracao = extrair("conversa.docx", _docx(TRANSCRICAO))
    assert extracao.tem_tempo is False
    assert any("latência" in aviso for aviso in extracao.avisos)


def test_transcricao_com_horario_e_medida_de_verdade():
    com_hora = (
        "[10:00] Cliente: bom dia\n"
        "[10:02] Bot: bom dia, como posso ajudar?\n"
    )
    extracao = extrair("conversa.docx", _docx(com_hora))
    assert extracao.tem_tempo is True
    mensagens = extracao.conversas[0].mensagens
    assert (mensagens[1].enviada_em - mensagens[0].enviada_em).total_seconds() == 120


def test_linha_sem_autor_continua_a_fala_anterior():
    """Word quebra paragrafo no meio da fala; picar viraria predicoes separadas."""
    texto = "Cliente: bom dia\nminha cobranca veio duplicada\nBot: vou verificar\n"
    extracao = extrair("conversa.docx", _docx(texto))
    mensagens = extracao.conversas[0].mensagens
    assert len(mensagens) == 2
    assert "duplicada" in mensagens[0].texto


# ---------------------------------------------------------------------------
# recusas: cada uma tem que NOMEAR o que se esperava


def test_extensao_desconhecida_lista_as_aceitas():
    with pytest.raises(ArquivoIlegivelError) as erro:
        extrair("foto.png", b"qualquer coisa")
    assert ".csv" in str(erro.value) and ".pdf" in str(erro.value)


def test_arquivo_sem_extensao_e_recusado_dizendo_o_motivo():
    with pytest.raises(ArquivoIlegivelError) as erro:
        extrair("arquivo", b"a,b\n1,2\n")
    assert "extensao" in str(erro.value)


def test_csv_com_colunas_erradas_diz_quais_faltam_e_quais_achou():
    with pytest.raises(ArquivoIlegivelError) as erro:
        extrair("errado.csv", b"nome,idade\nana,30\n")
    mensagem = str(erro.value)
    assert "conversa_id" in mensagem  # o que falta
    assert "nome" in mensagem  # o que achou


def test_transcricao_sem_fala_reconhecida_mostra_um_exemplo():
    """Recusar sem exemplo obriga a adivinhar o formato esperado."""
    with pytest.raises(ArquivoIlegivelError) as erro:
        extrair("vazio.docx", _docx("um texto qualquer sem estrutura\noutra linha\n"))
    assert "Cliente:" in str(erro.value)


def test_csv_vazio_e_recusado():
    with pytest.raises(ArquivoIlegivelError):
        extrair("vazio.csv", b"")


def test_pdf_sem_texto_extraivel_explica_que_nao_ha_ocr():
    """PDF escaneado e imagem: dizer "ilegivel" mandaria procurar o problema errado."""
    from pypdf import PdfWriter

    escritor = PdfWriter()
    escritor.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    escritor.write(buffer)

    with pytest.raises(ArquivoIlegivelError) as erro:
        extrair("scan.pdf", buffer.getvalue())
    assert "OCR" in str(erro.value)
